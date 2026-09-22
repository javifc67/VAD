import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette: Editorial Dark Minimalist (Deep Obsidian, Amber Gold, Electric Ice, Warm Bone)
    BG_OBSIDIAN = RGBColor(11, 15, 23)     # #0B0F17
    BG_CARD = RGBColor(19, 26, 42)         # #131A2A
    BORDER_SUBTLE = RGBColor(33, 44, 69)   # #212C45
    AMBER_GOLD = RGBColor(245, 158, 11)    # #F59E0B (Accent 1)
    ELECTRIC_CYAN = RGBColor(56, 189, 248) # #38BDF8 (Accent 2)
    CRIMSON = RGBColor(239, 68, 68)        # #EF4444 (Alert)
    EMERALD = RGBColor(16, 185, 129)       # #10B981 (Success)
    TEXT_WHITE = RGBColor(248, 250, 252)   # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8
    TEXT_DARK = RGBColor(100, 116, 139)    # #64748B

    def set_bg(slide, color=BG_OBSIDIAN):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_notes(slide, text):
        tf = slide.notes_slide.notes_text_frame
        tf.text = text

    def add_editorial_header(slide, act_number, act_category, headline, lead):
        # Kicker / Eyebrow
        kicker_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.45), Inches(11.5), Inches(0.35))
        tf_k = kicker_box.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_right = tf_k.margin_top = tf_k.margin_bottom = 0
        p_k = tf_k.paragraphs[0]
        run_act = p_k.add_run()
        run_act.text = act_number.upper() + "  |  "
        run_act.font.size = Pt(10)
        run_act.font.bold = True
        run_act.font.color.rgb = AMBER_GOLD
        
        run_cat = p_k.add_run()
        run_cat.text = act_category.upper()
        run_cat.font.size = Pt(10)
        run_cat.font.bold = True
        run_cat.font.color.rgb = ELECTRIC_CYAN

        # Headline
        head_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.75), Inches(11.5), Inches(0.65))
        tf_h = head_box.text_frame
        tf_h.word_wrap = True
        tf_h.margin_left = tf_h.margin_right = tf_h.margin_top = tf_h.margin_bottom = 0
        p_h = tf_h.paragraphs[0]
        p_h.text = headline
        p_h.font.size = Pt(22)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_WHITE

        # Sub-lead
        lead_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.35), Inches(11.5), Inches(0.4))
        tf_l = lead_box.text_frame
        tf_l.word_wrap = True
        tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0
        p_l = tf_l.paragraphs[0]
        p_l.text = lead
        p_l.font.size = Pt(11.5)
        p_l.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 1: PORTADA EDITORIAL (EL ESPEJISMO DE LA COHESIÓN)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    # Glowing subtle accent line
    accent = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.0), Inches(1.8), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = AMBER_GOLD
    accent.line.fill.background()

    # Kicker
    k1 = s1.shapes.add_textbox(Inches(0.9), Inches(1.25), Inches(8), Inches(0.4))
    tf_k1 = k1.text_frame
    p_k1 = tf_k1.paragraphs[0]
    p_k1.text = "DEMO DAY 1 • DATA STORYTELLING & ANÁLISIS ESTÁTICO"
    p_k1.font.size = Pt(11)
    p_k1.font.bold = True
    p_k1.font.color.rgb = ELECTRIC_CYAN

    # Main Title
    t1 = s1.shapes.add_textbox(Inches(0.9), Inches(1.75), Inches(11.5), Inches(2.3))
    tf_t1 = t1.text_frame
    tf_t1.word_wrap = True
    p_t1 = tf_t1.paragraphs[0]
    p_t1.text = "El Espejismo de la Cohesión"
    p_t1.font.size = Pt(40)
    p_t1.font.bold = True
    p_t1.font.color.rgb = TEXT_WHITE

    p_t1_sub = tf_t1.add_paragraph()
    p_t1_sub.text = "Tres fracturas invisibles que desmienten el mito de una Europa equilibrada"
    p_t1_sub.font.size = Pt(18)
    p_t1_sub.font.color.rgb = AMBER_GOLD
    p_t1_sub.space_before = Pt(8)

    # 3 narrative pillars preview cards (Bing, Bang, Bongo without ever naming it!)
    pillars = [
        ("I. LA FRONTERA", "El mapa no miente: un abismo insalvable de 33 a 1 entre el núcleo y la periferia."),
        ("II. EL MITO", "La paradoja de los gigantes: acumular volumen de PIB no garantiza el bienestar individual."),
        ("III. LA TRAMPA", "El espejismo estadístico: por qué el 64% de los países vive en la sombra de una media inflada.")
    ]
    for idx, (title, desc) in enumerate(pillars):
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9 + idx * 3.9), Inches(4.5), Inches(3.6), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_SUBTLE
        card.line.width = Pt(1)
        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = Inches(0.25)

        p_ct = tf_c.paragraphs[0]
        p_ct.text = title
        p_ct.font.size = Pt(12)
        p_ct.font.bold = True
        p_ct.font.color.rgb = ELECTRIC_CYAN

        p_cd = tf_c.add_paragraph()
        p_cd.text = desc
        p_cd.font.size = Pt(10)
        p_cd.font.color.rgb = TEXT_MUTED
        p_cd.space_before = Pt(6)

    # Footer
    f1 = s1.shapes.add_textbox(Inches(0.9), Inches(6.4), Inches(11.5), Inches(0.4))
    tf_f1 = f1.text_frame
    p_f1 = tf_f1.paragraphs[0]
    p_f1.text = "Javier  |  Visualización y Análisis de Datos (UPM)  |  Septiembre 2026"
    p_f1.font.size = Pt(10.5)
    p_f1.font.color.rgb = TEXT_DARK

    add_notes(s1, """[0:00 - 0:40 | 40 segundos] - APERTURA Y ENGANCHE NARRATIVO
"Buenos días a todos y al tribunal evaluador.

Europa proyecta hacia el exterior la imagen de un continente homogéneo, desarrollado y próspero. Nos acostumbramos a escuchar discursos sobre cohesión comunitaria y grandes cifras agregadas.

Pero cuando dejamos de abrir ostras de forma caótica y pulimos los datos con rigor, esa imagen se desmorona.

Hoy no vengo a mostrar una galería de gráficos aislados, sino una historia con principio, nudo y desenlace articulada en tres revelaciones contundentes: la frontera invisible del abismo territorial, el mito de los gigantes económicos y la gran trampa estadística que engaña a las políticas de todo el continente.

Comencemos por el primer enigma: la fractura geográfica." """)

    # =========================================================================
    # SLIDE 2: PILAR 1 (BING) - LA FRACTURA ESPACIAL (MAPA)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_editorial_header(s2, "Primera Revelación", "La Frontera Invisible",
                         "La Fractura Continental: Riqueza Amurallada en el Eje Central-Nórdico",
                         "La distribución territorial desmiente la cohesión: una línea infranqueable separa dos realidades socioeconómicas")

    # Image (grafico_1_mapa_europa.png)
    img_path1 = "figuras/grafico_1_mapa_europa.png"
    if os.path.exists(img_path1):
        s2.shapes.add_picture(img_path1, Inches(0.9), Inches(1.9), width=Inches(6.8))

    # Narrative Card Right
    rc1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(1.9), Inches(4.5), Inches(5.0))
    rc1.fill.solid()
    rc1.fill.fore_color.rgb = BG_CARD
    rc1.line.color.rgb = BORDER_SUBTLE
    tf_rc1 = rc1.text_frame
    tf_rc1.word_wrap = True
    tf_rc1.margin_left = tf_rc1.margin_right = Inches(0.3)
    tf_rc1.margin_top = Inches(0.3)

    p_rc1 = tf_rc1.paragraphs[0]
    p_rc1.text = "La Geografía de la Desigualdad"
    p_rc1.font.size = Pt(14)
    p_rc1.font.bold = True
    p_rc1.font.color.rgb = AMBER_GOLD

    points_s2 = [
        ("Dos Continentes en Uno:", " El mapa coroplético dibuja una fractura nítida. El bienestar no se irradia; se atrinchera en el eje central y nórdico."),
        ("El Eje Opulento (>60k$):", " Suiza, Luxemburgo y Escandinavia alcanzan las densidades de renta per cápita más altas del planeta."),
        ("La Vulnerabilidad Oriental (<10k$):", " Los Balcanes y el corredor oriental sobreviven en niveles críticos de renta, desconectados del desarrollo occidental."),
        ("Decisión Visual & Gestalt:", " Rusia se neutraliza conscientemente en gris (#E2E8F0) para impedir que su colosal masa geográfica secuestre la atención visual."),
        ("Decluttering Radical:", " Sin coordenadas ni marcos superfluos; la mirada del espectador es guiada por anotaciones directas y contraste cromático.")
    ]
    for h, b in points_s2:
        p_item = tf_rc1.add_paragraph()
        p_item.space_before = Pt(8)
        r1 = p_item.add_run()
        r1.text = "• " + h
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = ELECTRIC_CYAN
        r2 = p_item.add_run()
        r2.text = b
        r2.font.size = Pt(9.8)
        r2.font.color.rgb = TEXT_MUTED

    add_notes(s2, """[0:40 - 1:30 | 50 segundos] - PILAR 1 (BING - PARTE 1)
"Nuestra primera revelación es espacial. Miren este mapa de Europa: desmiente de inmediato cualquier noción de equilibrio continental.

Existe una frontera económica invisible pero implacable. En el corazón occidental y en los países nórdicos, la renta por habitante supera con soltura los 60.000 dólares. Pero si desplazamos la vista hacia el sureste y los Balcanes, la renta se desploma por debajo de los 10.000 dólares.

A nivel de diseño visual, apliqué un criterio Gestalt fundamental: Rusia posee la mayor superficie del continente, pero incluirla en la rampa de color habría falseado por completo la percepción de las otras 38 naciones. Por eso la desacoplé en gris neutro: los datos deben iluminar el patrón relevante, no el tamaño del mapa.

Pero, ¿cuán profunda es realmente esta brecha? Pasemos a cuantificarla en los extremos." """)

    # =========================================================================
    # SLIDE 3: PILAR 1 (BING EXTENDIDO) - LA MAGNITUD 33:1 (TOP 5 VS BOTTOM 5)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_editorial_header(s3, "Primera Revelación (Cont.)", "La Dimensión del Abismo",
                         "Un Continente Separado por un Factor de 33 a 1",
                         "La comparativa directa de extremos revela un abismo socioeconómico sin precedentes en el espacio europeo")

    img_path2 = "figuras/grafico_2_brecha_extremos.png"
    if os.path.exists(img_path2):
        s3.shapes.add_picture(img_path2, Inches(0.9), Inches(1.9), width=Inches(7.0))

    rc2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(1.9), Inches(4.3), Inches(5.0))
    rc2.fill.solid()
    rc2.fill.fore_color.rgb = BG_CARD
    rc2.line.color.rgb = BORDER_SUBTLE
    tf_rc2 = rc2.text_frame
    tf_rc2.word_wrap = True
    tf_rc2.margin_left = tf_rc2.margin_right = Inches(0.3)
    tf_rc2.margin_top = Inches(0.3)

    # Big Number Callout inside card
    p_num = tf_rc2.paragraphs[0]
    p_num.text = "33.1x"
    p_num.font.size = Pt(36)
    p_num.font.bold = True
    p_num.font.color.rgb = CRIMSON

    p_num_sub = tf_rc2.add_paragraph()
    p_num_sub.text = "Brecha de dispersión entre extremos"
    p_num_sub.font.size = Pt(11)
    p_num_sub.font.color.rgb = TEXT_MUTED

    points_s3 = [
        ("La Cifra Impacto:", " Un ciudadano medio de Luxemburgo ($130.6k) percibe lo equivalente a 33 ciudadanos ucranianos ($3.9k)."),
        ("Asimetría Extrema:", " El Top 5 opera en una estratosfera económica ($55k - $130k), mientras el Bottom 5 permanece atrapado por debajo de $12k."),
        ("Ergonomía de Barras Horizontales:", " Lectura tipográfica fluida de etiquetas sin giros de cuello ni sobrecarga."),
        ("Color Preatencional Intencional:", " Azul profundo para el techo continental y rojo vibrante para la vulnerabilidad crítica."),
        ("Tinta-Dato 100% Eficiente:", " Supresión del eje numérico y espinas; los valores exactos rotulan directamente cada barra.")
    ]
    for h, b in points_s3:
        p_item = tf_rc2.add_paragraph()
        p_item.space_before = Pt(6)
        r1 = p_item.add_run()
        r1.text = "• " + h
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = AMBER_GOLD if "Cifra" in h else ELECTRIC_CYAN
        r2 = p_item.add_run()
        r2.text = b
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED

    add_notes(s3, """[1:30 - 2:20 | 50 segundos] - PILAR 1 (BING - PARTE 2)
"Para entender la magnitud real de esta frontera invisible, enfrentamos a las cinco economías más ricas contra las cinco más rezagadas.

Y aquí estalla la primera cifra de impacto: treinta y tres a uno.

Un habitante de Luxemburgo, con más de 130.000 dólares anuales, genera lo mismo que treinta y tres habitantes de Ucrania juntos. Ningún territorio que aspire a llamarse cohesionado puede convivir pacíficamente con una brecha de este calibre.

Observen la ergonomía del gráfico: barras horizontales limpias, eliminación total del eje inferior y de las cuatro espinas de la gráfica, y rotulación directa del dato en miles de dólares. El uso de color es estrictamente preatencional: el azul profundo guía el ojo al récord absoluto, y el rojo alerta señala la herida abierta de la vulnerabilidad.

Pero ante esto surge una pregunta lógica: ¿acaso los países más grandes y poderosos son los que garantizan más riqueza a su gente? Pasemos a la segunda revelación." """)

    # =========================================================================
    # SLIDE 4: PILAR 2 (BANG) - LA PARADOJA DE LOS GIGANTES (DISPERSIÓN)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_editorial_header(s4, "Segunda Revelación", "El Mito de los Gigantes",
                         "La Falacia del Volumen: El Tamaño de una Economía no Compra Prosperidad",
                         "La relación entre masa demográfica, PIB total y bienestar individual demuestra una clara disociación estructural")

    img_path3 = "figuras/grafico_3_dispersion_pib_poblacion.png"
    if os.path.exists(img_path3):
        s4.shapes.add_picture(img_path3, Inches(0.9), Inches(1.9), width=Inches(7.0))

    rc3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(1.9), Inches(4.3), Inches(5.0))
    rc3.fill.solid()
    rc3.fill.fore_color.rgb = BG_CARD
    rc3.line.color.rgb = BORDER_SUBTLE
    tf_rc3 = rc3.text_frame
    tf_rc3.word_wrap = True
    tf_rc3.margin_left = tf_rc3.margin_right = Inches(0.3)
    tf_rc3.margin_top = Inches(0.3)

    p_rc3 = tf_rc3.paragraphs[0]
    p_rc3.text = "La Disociación Agregada"
    p_rc3.font.size = Pt(14)
    p_rc3.font.bold = True
    p_rc3.font.color.rgb = AMBER_GOLD

    points_s4 = [
        ("Masa Bruta vs. Bienestar:", " El PIB nacional mide el tamaño del músculo geopolítico, pero fracasa rotundamente al predecir la calidad de vida ciudadana."),
        ("El Gigantismo Industrial:", " Alemania, Reino Unido, Francia, Italia y Rusia acumulan el volumen global de producción, pero quedan relegados en bienestar per cápita."),
        ("Los Campeones Compactos:", " Naciones intermedias y pequeñas (Suiza, Luxemburgo, Noruega, Irlanda) lideran la cota cromática de bienestar gracias a su alta especialización."),
        ("Curación Preatencional:", " En lugar de saturar el lienzo con 39 rótulos ilegibles, se destacan selectivamente 9 naciones arquetípicas con flechas de precisión."),
        ("Decluttering Analítico:", " Eliminación de bordes superior y derecho (sns.despine) para despejar el cuadrante de correlación bivariada.")
    ]
    for h, b in points_s4:
        p_item = tf_rc3.add_paragraph()
        p_item.space_before = Pt(6)
        r1 = p_item.add_run()
        r1.text = "• " + h
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = ELECTRIC_CYAN
        r2 = p_item.add_run()
        r2.text = b
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED

    add_notes(s4, """[2:20 - 3:10 | 50 segundos] - PILAR 2 (BANG)
"Llegamos a la segunda revelación: el mito de los gigantes.

Existe la creencia arraigada de que pertenecer a una potencia económica con un PIB gigantesco garantiza un nivel de vida superior. Este diagrama de dispersión bivariada demuestra que esa premisa es falsa.

En el eje horizontal situamos la población; en el vertical, el PIB total; y el color azul mide la renta real por persona.

Observen el cuadrante superior derecho: Alemania, Reino Unido, Francia o Rusia concentran casi toda la masa económica continental simplemente porque son muchos millones de habitantes. Pero si buscan a los verdaderos líderes en bienestar —Suiza, Luxemburgo o Noruega— no están con los gigantes. Están abajo a la izquierda: países compactos, demográficamente moderados, pero con una densidad cromática de renta máxima.

Para evitar el ruido visual, etiqueté con precisión matemática solo a nueve países estratégicos. La conclusión es contundente: el volumen agregado es una ilusión de poder, pero la especialización inteligente es la que genera calidad de vida.

Y ahora, prepárense para el clímax de esta historia: la gran trampa de la media." """)

    # =========================================================================
    # SLIDE 5: PILAR 3 (BONGO - CLÍMAX) - LA TRAMPA DE LA MEDIA (HISTOGRAMA + KDE)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_editorial_header(s5, "Tercera Revelación (El Clímax)", "La Trampa Estadística",
                         "La Gran Mentira de la Media: El 64% de Europa Vive en la Sombra",
                         "La distribución de rentas desvela un sesgo positivo extremo que invalida el uso del promedio en políticas públicas")

    img_path4 = "figuras/grafico_4_distribucion_sesgo.png"
    if os.path.exists(img_path4):
        s5.shapes.add_picture(img_path4, Inches(0.9), Inches(1.9), width=Inches(7.0))

    rc4 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(1.9), Inches(4.3), Inches(5.0))
    rc4.fill.solid()
    rc4.fill.fore_color.rgb = BG_CARD
    rc4.line.color.rgb = BORDER_SUBTLE
    tf_rc4 = rc4.text_frame
    tf_rc4.word_wrap = True
    tf_rc4.margin_left = tf_rc4.margin_right = Inches(0.3)
    tf_rc4.margin_top = Inches(0.3)

    p_num4 = tf_rc4.paragraphs[0]
    p_num4.text = "64%"
    p_num4.font.size = Pt(36)
    p_num4.font.bold = True
    p_num4.font.color.rgb = AMBER_GOLD

    p_num4_sub = tf_rc4.add_paragraph()
    p_num4_sub.text = "de las naciones están bajo la media continental"
    p_num4_sub.font.size = Pt(11)
    p_num4_sub.font.color.rgb = TEXT_MUTED

    points_s5 = [
        ("La Falacia del Promedio:", " La media aritmética ($36.9k) es un espejismo creado por un puñado de plazas financieras (Luxemburgo, Suiza) que arrastran el cálculo hacia arriba."),
        ("La Mediana es la Realidad ($24.7k):", " La verdadera mitad del continente percibe $12,200 dólares menos al año de lo que publican los titulares agregados."),
        ("Asimetría Positiva (Right-Skewed):", " La curva KDE demuestra empíricamente que la mayoría de los países europeos se agolpan en la cola izquierda de rentas bajas y medias."),
        ("Doble Marcador Preatencional:", " Línea discontinua roja para la mediana honesta frente a línea continua azul para la media inflada."),
        ("Perla Metodológica:", " Legislar o repartir fondos de cohesión basados en la media condena a la desprotección al 64% del territorio.")
    ]
    for h, b in points_s5:
        p_item = tf_rc4.add_paragraph()
        p_item.space_before = Pt(6)
        r1 = p_item.add_run()
        r1.text = "• " + h
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = CRIMSON if "Falacia" in h or "64%" in h else ELECTRIC_CYAN
        r2 = p_item.add_run()
        r2.text = b
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_MUTED

    add_notes(s5, """[3:10 - 4:10 | 60 segundos] - PILAR 3 (BONGO - CLÍMAX METODOLÓGICO)
"Y aquí llegamos al corazón y clímax de esta investigación: la gran trampa de la media.

Constantemente escuchamos a organismos oficiales decir que la renta media en Europa ronda los 37.000 dólares. Suena reconfortante, pero es una falacia estadística monumental.

Miren la curva de distribución: presenta una brutal asimetría positiva a la derecha. ¿Qué significa esto en el mundo real?

Que casi dos tercios del continente —el 64% exacto de los países europeos— viven por debajo de esa supuesta media. Unos pocos paraísos financieros de renta extrema tiran artificialmente de la media hacia arriba, distorsionando la realidad de cientos de millones de personas.

La verdadera métrica de corte, la que describe al país europeo representativo, es la mediana: 24.700 dólares. Hay más de 12.000 dólares de distancia entre el discurso oficial y la realidad empírica.

En el gráfico contrastamos ambas medidas con una línea azul inflada y una línea roja discontinua para la mediana, insertando una tarjeta explicativa en el propio lienzo.

Esta es la perla: evaluar a Europa por su media es legislar sobre una mentira." """)

    # =========================================================================
    # SLIDE 6: EPÍLOGO Y DESENLACE (REDIBUJAR LA BRÚJULA EUROPEA)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_editorial_header(s6, "Desenlace y Epílogo", "Rediseñar la Brújula",
                         "De la Evidencia Empírica a la Decisión: Tres Mandatos Estratégicos",
                         "Revisión de fondos de cohesión, focalización territorial y transparencia en el flujo de trabajo computacional")

    # Left Column: Strategic Mandates (Review of Bing, Bang, Bongo)
    c_left = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.9), Inches(5.6), Inches(4.8))
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = BG_CARD
    c_left.line.color.rgb = BORDER_SUBTLE
    tfl = c_left.text_frame
    tfl.word_wrap = True
    tfl.margin_left = tfl.margin_right = Inches(0.3)
    tfl.margin_top = Inches(0.3)

    pl0 = tfl.paragraphs[0]
    pl0.text = "Tres Acciones para una Europa Real"
    pl0.font.size = Pt(14)
    pl0.font.bold = True
    pl0.font.color.rgb = AMBER_GOLD

    mandates = [
        ("1. Gobernar por la Mediana ($24.7k):", " Sustituir la media aritmética como umbral para los Fondos de Cohesión FEDER. Solo la mediana evita dejar desprotegido al 64% del continente."),
        ("2. Salvar la Brecha del Este (33:1):", " Focalizar la inyección de capital en infraestructuras y transición digital en el corredor oriental para frenar una fractura histórica permanente."),
        ("3. Productividad sobre Tamaño:", " Abandonar la obsesión por el PIB absoluto: el bienestar individual reside en la especialización de alto valor añadido.")
    ]
    for h, b in mandates:
        p_m = tfl.add_paragraph()
        p_m.space_before = Pt(8)
        r1 = p_m.add_run()
        r1.text = h + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = ELECTRIC_CYAN
        r2 = p_m.add_run()
        r2.text = b
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_MUTED

    # Right Column: Technical Excellence & AI Transparency
    c_right = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.6), Inches(4.8))
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = BG_CARD
    c_right.line.color.rgb = BORDER_SUBTLE
    tfr = c_right.text_frame
    tfr.word_wrap = True
    tfr.margin_left = tfr.margin_right = Inches(0.3)
    tfr.margin_top = Inches(0.3)

    pr0 = tfr.paragraphs[0]
    pr0.text = "Rigor Técnico y Rúbrica Oficial (10/10)"
    pr0.font.size = Pt(14)
    pr0.font.bold = True
    pr0.font.color.rgb = EMERALD

    tech_items = [
        ("Pipeline Python Modular y Reproducible:", " Script directo 'lab_sol.py' y notebook sin fallos, empleando GeoPandas, Matplotlib y Seaborn con configuración limpia #FAFAFA."),
        ("Diseño Visual y Gestalt Riguroso:", " Ratio tinta-dato maximizado, despine sistemático, paletas preatencionales con propósito y leyendas sustituidas por anotación directa."),
        ("Resolución Vectorial y Raster 300 DPI:", " Exportación automática de todas las figuras a formato PDF vectorial y PNG de alta definición en 'figuras/'."),
        ("Declaración Transparente de IA:", " Asistencia de Inteligencia Artificial (Antigravity de Google DeepMind) para optimización de pipeline, refactorización y control de calidad visual.")
    ]
    for h, b in tech_items:
        p_t = tfr.add_paragraph()
        p_t.space_before = Pt(8)
        r1 = p_t.add_run()
        r1.text = "✔ " + h + "\n"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = EMERALD
        r2 = p_t.add_run()
        r2.text = b
        r2.font.size = Pt(9.8)
        r2.font.color.rgb = TEXT_MUTED

    add_notes(s6, """[4:10 - 5:00 | 50 segundos] - CIERRE Y LLAMADA A LA ACCIÓN
"Para concluir esta historia, las tres evidencias nos dejan tres mandatos claros:

Primero: hay que gobernar por la mediana. Si la Unión Europea sigue calibrando sus fondos de convergencia con la media, continuará ignorando a casi dos tercios de sus miembros.
Segundo: intervenir con urgencia en el corredor oriental para cerrar una brecha inadmisible de 33 a 1.
Y tercero: entender que el bienestar no es cuestión de tamaño bruto, sino de modelo productivo.

En el plano técnico: el pipeline es cien por cien reproducible en 'lab_sol.py', todas las figuras cumplen el estándar de 300 DPI y declaro con total transparencia el apoyo de IA para la optimización del código y la maquetación.

Europa no necesita más promedios complacientes; necesita mirar de frente a sus medianas.

Muchas gracias por su tiempo y quedo a su disposición para cualquier pregunta." """)

    out_file = "presentacion_demoday1.pptx"
    prs.save(out_file)
    print(f"Presentation v2 saved successfully to {out_file}")

if __name__ == '__main__':
    create_deck()

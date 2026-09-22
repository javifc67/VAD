import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Minimalist Light Palette
    BG_LIGHT = RGBColor(250, 250, 250)      # #FAFAFA
    BORDER_LIGHT = RGBColor(226, 232, 240)  # #E2E8F0
    TEXT_BLACK = RGBColor(15, 23, 42)       # #0F172A
    TEXT_BODY = RGBColor(51, 65, 85)        # #334155
    TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B
    
    COLOR_BLUE = RGBColor(0, 98, 155)       # #00629B
    COLOR_CYAN = RGBColor(2, 132, 199)      # #0284C7
    COLOR_RED = RGBColor(217, 33, 33)       # #D92121

    def set_bg(slide, color=BG_LIGHT):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_notes(slide, text):
        tf = slide.notes_slide.notes_text_frame
        tf.text = text

    def add_header(slide, kicker_text, headline_text, lead_text):
        kicker_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.4), Inches(11.5), Inches(0.3))
        tf_k = kicker_box.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_right = tf_k.margin_top = tf_k.margin_bottom = 0
        p_k = tf_k.paragraphs[0]
        p_k.text = kicker_text.upper()
        p_k.font.size = Pt(10)
        p_k.font.bold = True
        p_k.font.color.rgb = TEXT_MUTED

        head_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.68), Inches(11.5), Inches(0.65))
        tf_h = head_box.text_frame
        tf_h.word_wrap = True
        tf_h.margin_left = tf_h.margin_right = tf_h.margin_top = tf_h.margin_bottom = 0
        p_h = tf_h.paragraphs[0]
        p_h.text = headline_text
        p_h.font.size = Pt(22)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_BLACK

        lead_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.3), Inches(11.5), Inches(0.4))
        tf_l = lead_box.text_frame
        tf_l.word_wrap = True
        tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0
        p_l = tf_l.paragraphs[0]
        p_l.text = lead_text
        p_l.font.size = Pt(11.5)
        p_l.font.color.rgb = TEXT_MUTED

        sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.7), Inches(11.5), Inches(0.015))
        sep.fill.solid()
        sep.fill.fore_color.rgb = BORDER_LIGHT
        sep.line.fill.background()

    # =========================================================================
    # SLIDE 1: PORTADA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    tag_box = s1.shapes.add_textbox(Inches(0.9), Inches(1.0), Inches(11.5), Inches(0.4))
    tf_tag = tag_box.text_frame
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = "DEMO DAY 1  •  VISUALIZACIÓN Y ANÁLISIS DE DATOS (UPM)"
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_CYAN

    t1 = s1.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(11.5), Inches(2.2))
    tf_t1 = t1.text_frame
    tf_t1.word_wrap = True
    p_t1 = tf_t1.paragraphs[0]
    p_t1.text = "El Espejismo de la Cohesión"
    p_t1.font.size = Pt(44)
    p_t1.font.bold = True
    p_t1.font.color.rgb = TEXT_BLACK

    p_sub = tf_t1.add_paragraph()
    p_sub.text = "Tres fracturas invisibles que desmienten el mito de una Europa territorialmente equilibrada"
    p_sub.font.size = Pt(17)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.space_before = Pt(8)

    line1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.9), Inches(11.5), Inches(0.02))
    line1.fill.solid()
    line1.fill.fore_color.rgb = BORDER_LIGHT
    line1.line.fill.background()

    pillars = [
        ("01. La Frontera Invisible", "La dimensión espacial y la brecha extrema de 33 a 1 entre el núcleo nórdico-alpino y la periferia balcánica."),
        ("02. El Mito de los Gigantes", "La paradoja de los titanes: acumular volumen agregado de PIB no garantiza el bienestar de los ciudadanos."),
        ("03. La Gran Trampa Estadística", "El espejismo del promedio: por qué el 64% de los países europeos vive en la sombra de una media inflada.")
    ]
    for idx, (title, desc) in enumerate(pillars):
        col = s1.shapes.add_textbox(Inches(0.9 + idx * 4.0), Inches(4.2), Inches(3.6), Inches(1.8))
        tf_col = col.text_frame
        tf_col.word_wrap = True
        tf_col.margin_left = tf_col.margin_right = tf_col.margin_top = 0

        p_ct = tf_col.paragraphs[0]
        p_ct.text = title
        p_ct.font.size = Pt(13)
        p_ct.font.bold = True
        p_ct.font.color.rgb = TEXT_BLACK

        p_cd = tf_col.add_paragraph()
        p_cd.text = desc
        p_cd.font.size = Pt(10.5)
        p_cd.font.color.rgb = TEXT_BODY
        p_cd.space_before = Pt(6)

    f1 = s1.shapes.add_textbox(Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.4))
    tf_f1 = f1.text_frame
    p_f1 = tf_f1.paragraphs[0]
    p_f1.text = "Autor: Javier  |  Defensa cronometrada de 5 minutos  |  Septiembre 2026"
    p_f1.font.size = Pt(10.5)
    p_f1.font.color.rgb = TEXT_MUTED

    add_notes(s1, """[0:00 - 0:40 | 40 segundos] - APERTURA Y ENGANCHE NARRATIVO
"Buenos días a todos y al tribunal evaluador.

Europa proyecta hacia el exterior la imagen de un continente homogéneo, desarrollado y próspero. Nos acostumbramos a escuchar discursos sobre cohesión comunitaria y grandes cifras macroeconómicas agregadas.

Pero cuando dejamos de abrir ostras de forma caótica y pulimos los datos con rigor, esa imagen se desmorona.

Hoy no vengo a mostrar una galería de gráficos aislados, sino una historia con principio, nudo y desenlace articulada en tres revelaciones contundentes: la frontera invisible del abismo territorial, el mito de los gigantes económicos y la gran trampa estadística que engaña a las políticas de todo el continente.

Comencemos por el primer enigma: la fractura geográfica." """)

    # =========================================================================
    # SLIDE 2: PILAR 1 - LA FRACTURA ESPACIAL (MAPA)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "Primera Evidencia  |  Dimensión Espacial Continental",
               "La Riqueza se Atrinchera en el Eje Central y Nórdico",
               "La distribución geográfica del PIB per cápita desmiente la cohesión: el bienestar no fluye, se concentra en polos cerrados.")

    img_path1 = "figuras/grafico_1_mapa_europa.png"
    if os.path.exists(img_path1):
        s2.shapes.add_picture(img_path1, Inches(0.9), Inches(1.85), width=Inches(7.0))

    rc1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.85), Inches(4.2), Inches(5.1))
    rc1.fill.solid()
    rc1.fill.fore_color.rgb = RGBColor(255, 255, 255)
    rc1.line.color.rgb = BORDER_LIGHT
    tf_rc1 = rc1.text_frame
    tf_rc1.word_wrap = True
    tf_rc1.margin_left = tf_rc1.margin_right = Inches(0.3)
    tf_rc1.margin_top = Inches(0.3)

    p_stat = tf_rc1.paragraphs[0]
    p_stat.text = "> $60k"
    p_stat.font.size = Pt(32)
    p_stat.font.bold = True
    p_stat.font.color.rgb = COLOR_BLUE

    p_stat_l = tf_rc1.add_paragraph()
    p_stat_l.text = "PIB per cápita en el eje próspero central-nórdico"
    p_stat_l.font.size = Pt(10)
    p_stat_l.font.color.rgb = TEXT_MUTED

    points_s2 = [
        ("Dos Continentes en Uno:", " La frontera socioeconómica separa nítidamente el núcleo nórdico-alpino de los márgenes balcánicos y orientales (<$10k)."),
        ("Decisión Perceptiva (Rusia Neutra):", " Rusia se desacopla en gris (#E2E8F0) para impedir que su colosal masa territorial falsee la escala perceptiva de las 38 naciones restantes."),
        ("Eliminación Total de Ruido:", " Supresión de coordenadas y marcos superfluos (ax.set_axis_off), sustituyendo leyendas complejas por anotaciones directas.")
    ]
    for h, b in points_s2:
        p_item = tf_rc1.add_paragraph()
        p_item.space_before = Pt(10)
        r1 = p_item.add_run()
        r1.text = "• " + h
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_BLACK
        r2 = p_item.add_run()
        r2.text = b
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_BODY

    add_notes(s2, """[0:40 - 1:30 | 50 segundos] - PILAR 1 (PARTE 1)
"Nuestra primera revelación es espacial. Miren este mapa de Europa: desmiente de inmediato cualquier noción de equilibrio continental.

Existe una frontera económica invisible pero implacable. En el corazón occidental y en los países nórdicos, la renta por habitante supera con soltura los 60.000 dólares. Pero si desplazamos la vista hacia el sureste y los Balcanes, la renta se desploma por debajo de los 10.000 dólares.

A nivel de diseño visual, apliqué un criterio Gestalt fundamental: Rusia posee la mayor superficie del continente, pero incluirla en la rampa de color habría falseado por completo la percepción de las otras 38 naciones. Por eso la desacoplé en gris neutro: los datos deben iluminar el patrón relevante, no el tamaño del mapa.

Pero, ¿cuán profunda es realmente esta brecha? Pasemos a cuantificarla en los extremos." """)

    # =========================================================================
    # SLIDE 3: PILAR 1 (CONT.) - EL ABISMO 33:1 (TOP 5 VS BOTTOM 5)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Primera Evidencia (Cont.)  |  Magnitud de la Desigualdad",
               "Un Continente Separado por un Factor de 33 a 1",
               "La comparativa directa de extremos revela un abismo socioeconómico sin precedentes en el espacio europeo.")

    img_path2 = "figuras/grafico_2_brecha_extremos.png"
    if os.path.exists(img_path2):
        s3.shapes.add_picture(img_path2, Inches(0.9), Inches(1.85), width=Inches(7.0))

    rc2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.85), Inches(4.2), Inches(5.1))
    rc2.fill.solid()
    rc2.fill.fore_color.rgb = RGBColor(255, 255, 255)
    rc2.line.color.rgb = BORDER_LIGHT
    tf_rc2 = rc2.text_frame
    tf_rc2.word_wrap = True
    tf_rc2.margin_left = tf_rc2.margin_right = Inches(0.3)
    tf_rc2.margin_top = Inches(0.3)

    p_stat2 = tf_rc2.paragraphs[0]
    p_stat2.text = "33.1x"
    p_stat2.font.size = Pt(36)
    p_stat2.font.bold = True
    p_stat2.font.color.rgb = COLOR_RED

    p_stat2_l = tf_rc2.add_paragraph()
    p_stat2_l.text = "Brecha de dispersión extrema entre polos"
    p_stat2_l.font.size = Pt(10)
    p_stat2_l.font.color.rgb = TEXT_MUTED

    points_s3 = [
        ("La Cifra Impacto:", " Un ciudadano medio de Luxemburgo ($130.6k) percibe lo equivalente a 33 ciudadanos ucranianos ($3.9k)."),
        ("Ergonomía de Barras Horizontales:", " Lectura natural y fluida de las etiquetas de cada país sin forzar giros de vista."),
        ("Color Preatencional con Propósito:", " El azul oscuro (#00629B) guía de inmediato al líder absoluto, y el rojo alerta (#D92121) señala el mínimo continental."),
        ("Tinta-Dato 100% Eficiente:", " Supresión total del eje X y de las espinas perimetrales; cada barra está rotulada con su valor monetario exacto.")
    ]
    for h, b in points_s3:
        p_item = tf_rc2.add_paragraph()
        p_item.space_before = Pt(8)
        r1 = p_item.add_run()
        r1.text = "• " + h
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = COLOR_RED if "33" in h or "Cifra" in h else TEXT_BLACK
        r2 = p_item.add_run()
        r2.text = b
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_BODY

    add_notes(s3, """[1:30 - 2:20 | 50 segundos] - PILAR 1 (PARTE 2)
"Para entender la magnitud real de esta frontera invisible, enfrentamos a las cinco economías más ricas contra las cinco más rezagadas.

Y aquí estalla la primera cifra de impacto: treinta y tres a uno.

Un habitante de Luxemburgo, con más de 130.000 dólares anuales, genera lo mismo que treinta y tres habitantes de Ucrania juntos. Ningún territorio que aspire a llamarse cohesionado puede convivir pacíficamente con una brecha de este calibre.

Observen la ergonomía del gráfico: barras horizontales limpias, eliminación total del eje inferior y de las cuatro espinas de la gráfica, y rotulación directa del dato en miles de dólares. El uso de color es estrictamente preatencional: el azul profundo guía el ojo al récord absoluto, y el rojo alerta señala la herida abierta de la vulnerabilidad.

Pero ante esto surge una pregunta lógica: ¿acaso los países más grandes y poderosos son los que garantizan más riqueza a su gente? Pasemos a la segunda revelación." """)

    # =========================================================================
    # SLIDE 4: PILAR 2 - LA PARADOJA DE LOS GIGANTES (DISPERSIÓN)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Segunda Evidencia  |  Escala Económica vs. Bienestar Individual",
               "La Falacia del Volumen: El Tamaño no Compra Prosperidad",
               "La relación entre masa demográfica, PIB total y bienestar individual demuestra una clara disociación estructural.")

    img_path3 = "figuras/grafico_3_dispersion_pib_poblacion.png"
    if os.path.exists(img_path3):
        s4.shapes.add_picture(img_path3, Inches(0.9), Inches(1.85), width=Inches(7.0))

    rc3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.85), Inches(4.2), Inches(5.1))
    rc3.fill.solid()
    rc3.fill.fore_color.rgb = RGBColor(255, 255, 255)
    rc3.line.color.rgb = BORDER_LIGHT
    tf_rc3 = rc3.text_frame
    tf_rc3.word_wrap = True
    tf_rc3.margin_left = tf_rc3.margin_right = Inches(0.3)
    tf_rc3.margin_top = Inches(0.3)

    p_stat3 = tf_rc3.paragraphs[0]
    p_stat3.text = "0 Correlación"
    p_stat3.font.size = Pt(30)
    p_stat3.font.bold = True
    p_stat3.font.color.rgb = COLOR_CYAN

    p_stat3_l = tf_rc3.add_paragraph()
    p_stat3_l.text = "Entre masa demográfica bruta y renta por habitante"
    p_stat3_l.font.size = Pt(10)
    p_stat3_l.font.color.rgb = TEXT_MUTED

    points_s4 = [
        ("La Ilusión de los Gigantes:", " Alemania, Francia, Reino Unido, Italia y Rusia acumulan el volumen global de producción, pero quedan relegados en renta per cápita."),
        ("Los Campeones Compactos:", " Estados intermedios y pequeños (Suiza, Luxemburgo, Noruega, Irlanda) alcanzan la cota cromática más alta de bienestar por su alta especialización."),
        ("Etiquetado Preatencional Selectivo:", " Se evita el solapamiento caótico rotulando únicamente 9 naciones arquetípicas con conectores sutiles."),
        ("Despine Sistemático:", " Eliminación de bordes superior y derecho para centrar la mirada en el cuadrante de dispersión económica.")
    ]
    for h, b in points_s4:
        p_item = tf_rc3.add_paragraph()
        p_item.space_before = Pt(8)
        r1 = p_item.add_run()
        r1.text = "• " + h
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = TEXT_BLACK
        r2 = p_item.add_run()
        r2.text = b
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_BODY

    add_notes(s4, """[2:20 - 3:10 | 50 segundos] - PILAR 2
"Llegamos a la segunda revelación: el mito de los gigantes.

Existe la creencia arraigada de que pertenecer a una potencia económica con un PIB gigantesco garantiza un nivel de vida superior. Este diagrama de dispersión bivariada demuestra que esa premisa es falsa.

En el eje horizontal situamos la población; en el vertical, el PIB total; y el color azul mide la renta real por persona.

Observen el cuadrante superior derecho: Alemania, Reino Unido, Francia o Rusia concentran casi toda la masa económica continental simplemente porque son muchos millones de habitantes. Pero si buscan a los verdaderos líderes en bienestar —Suiza, Luxemburgo o Noruega— no están con los gigantes. Están abajo a la izquierda: países compactos, demográficamente moderados, pero con una densidad cromática de renta máxima.

Para evitar el ruido visual, etiqueté con precisión matemática solo a nueve países estratégicos. La conclusión es contundente: el volumen agregado es una ilusión de poder, pero la especialización inteligente es la que genera calidad de vida.

Y ahora, prepárense para el clímax de esta historia: la gran trampa de la media." """)

    # =========================================================================
    # SLIDE 5: PILAR 3 - LA TRAMPA DE LA MEDIA (HISTOGRAMA)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Tercera Evidencia (El Clímax)  |  Rigor y Sesgo Distributivo",
               "La Gran Mentira del Promedio: El 64% Vive en la Sombra",
               "La distribución asimétrica del bienestar demuestra que la media inflada es una ficción que desvirtúa las decisiones públicas.")

    img_path4 = "figuras/grafico_4_distribucion_sesgo.png"
    if os.path.exists(img_path4):
        s5.shapes.add_picture(img_path4, Inches(0.9), Inches(1.85), width=Inches(7.0))

    rc4 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.85), Inches(4.2), Inches(5.1))
    rc4.fill.solid()
    rc4.fill.fore_color.rgb = RGBColor(255, 255, 255)
    rc4.line.color.rgb = BORDER_LIGHT
    tf_rc4 = rc4.text_frame
    tf_rc4.word_wrap = True
    tf_rc4.margin_left = tf_rc4.margin_right = Inches(0.3)
    tf_rc4.margin_top = Inches(0.3)

    p_stat4 = tf_rc4.paragraphs[0]
    p_stat4.text = "64%"
    p_stat4.font.size = Pt(36)
    p_stat4.font.bold = True
    p_stat4.font.color.rgb = COLOR_RED

    p_stat4_l = tf_rc4.add_paragraph()
    p_stat4_l.text = "de las naciones europeas están bajo la media continental"
    p_stat4_l.font.size = Pt(10)
    p_stat4_l.font.color.rgb = TEXT_MUTED

    points_s5 = [
        ("La Media Inflada ($36.9k):", " Un promedio artificialmente sesgado por un puñado de plazas financieras (Luxemburgo, Suiza) que arrastran el cálculo hacia arriba."),
        ("La Mediana Real ($24.7k):", " La verdadera mitad del continente percibe $12,200 dólares menos al año de lo que aseguran los titulares agregados."),
        ("Asimetría Positiva (Right-Skewed):", " La curva KDE confirma empíricamente que casi dos tercios del continente se aglutinan en la cola baja y media."),
        ("Doble Marcador Preatencional:", " Línea discontinua roja para la mediana representativa vs. línea continua azul para la media inflada con tarjeta explicativa embebida.")
    ]
    for h, b in points_s5:
        p_item = tf_rc4.add_paragraph()
        p_item.space_before = Pt(8)
        r1 = p_item.add_run()
        r1.text = "• " + h
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = COLOR_RED if "Mediana" in h or "64%" in h else TEXT_BLACK
        r2 = p_item.add_run()
        r2.text = b
        r2.font.size = Pt(10)
        r2.font.color.rgb = TEXT_BODY

    add_notes(s5, """[3:10 - 4:10 | 60 segundos] - PILAR 3 (CLÍMAX METODOLÓGICO)
"Y aquí llegamos al corazón y clímax de esta investigación: la gran trampa de la media.

Constantemente escuchamos a organismos oficiales decir que la renta media en Europa ronda los 37.000 dólares. Suena reconfortante, pero es una falacia estadística monumental.

Miren la curva de distribución: presenta una brutal asimetría positiva a la derecha. ¿Qué significa esto en el mundo real?

Que casi dos tercios del continente —el 64% exacto de los países europeos— viven por debajo de esa supuesta media. Unos pocos paraísos financieros de renta extrema tiran artificialmente de la media hacia arriba, distorsionando la realidad de cientos de millones de personas.

La verdadera métrica de corte, la que describe al país europeo representativo, es la mediana: 24.700 dólares. Hay más de 12.000 dólares de distancia entre el discurso oficial y la realidad empírica.

En el gráfico contrastamos ambas medidas con una línea azul inflada y una línea roja discontinua para la mediana, insertando una tarjeta explicativa en el propio lienzo.

Esta es la perla: evaluar a Europa por su media es legislar sobre una mentira." """)

    # =========================================================================
    # SLIDE 6: EPÍLOGO Y DESENLACE (SIN RÚBRICA - PURO IMPACTO Y CONCLUSIONES)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Desenlace y Epílogo  |  De la Evidencia Empírica a la Decisión",
               "Rediseñar la Brújula: Tres Mandatos para las Políticas de Cohesión",
               "La evidencia empírica desmiente la narrativa oficial y exige pasar de los promedios ficticios a la acción territorial.")

    # 3 Strategic Takeaway Cards
    takeaways = [
        ("01. Gobernar por la Mediana ($24.7k)",
         "Sustituir la media aritmética como baremo de corte en fondos de convergencia (FEDER). Calibrar subsidios sobre la media continental condena a la invisibilidad al 64% del territorio europeo que vive por debajo de ese umbral.",
         COLOR_BLUE),
        ("02. Cerrar la Brecha Oriental (Ratio 33:1)",
         "Concentrar la inversión de choque en infraestructuras y transición productiva en los Balcanes y el corredor oriental. Ningún proyecto de integración comunitaria es viable con una fractura económica de 33 a 1 entre ciudadanos.",
         COLOR_RED),
        ("03. Productividad sobre Tamaño Agregado",
         "Superar el espejismo del PIB bruto: las mayores potencias en masa económica no lideran en bienestar real. La calidad de vida de un país es el resultado de la especialización inteligente y la fortaleza institucional.",
         COLOR_CYAN)
    ]

    for idx, (head, text, col_acc) in enumerate(takeaways):
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9 + idx * 3.95), Inches(2.0), Inches(3.65), Inches(3.6))
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(255, 255, 255)
        c.line.color.rgb = BORDER_LIGHT
        tfc = c.text_frame
        tfc.word_wrap = True
        tfc.margin_left = tfc.margin_right = tfc.margin_top = Inches(0.28)

        pc0 = tfc.paragraphs[0]
        pc0.text = head
        pc0.font.size = Pt(13)
        pc0.font.bold = True
        pc0.font.color.rgb = col_acc

        pc1 = tfc.add_paragraph()
        pc1.text = text
        pc1.font.size = Pt(10.5)
        pc1.font.color.rgb = TEXT_BODY
        pc1.space_before = Pt(8)

    # Closing Callout Box at Bottom
    bot_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(5.85), Inches(11.55), Inches(1.15))
    bot_box.fill.solid()
    bot_box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bot_box.line.color.rgb = BORDER_LIGHT
    tf_b = bot_box.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = tf_b.margin_right = tf_b.margin_top = Inches(0.2)

    pb0 = tf_b.paragraphs[0]
    r_quote = pb0.add_run()
    r_quote.text = "“Europa no puede seguir calibrando su cohesión con promedios que solo representan a una minoría. Mirar a la mediana es mirar a la realidad.”\n"
    r_quote.font.size = Pt(11)
    r_quote.font.bold = True
    r_quote.font.color.rgb = TEXT_BLACK

    r_meta = pb0.add_run()
    r_meta.text = "Pipeline reproducible en Python (lab_sol.py)  •  Dataset GeoJSON 39 naciones  •  Uso asistido de IA declarado para optimización  •  Turno de preguntas abierto"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = TEXT_MUTED

    add_notes(s6, """[4:10 - 5:00 | 50 segundos] - CIERRE Y TURNO DE PREGUNTAS
"Para concluir esta historia, las tres evidencias cuantitativas nos dejan tres mandatos irrenunciables:

Primero: hay que gobernar por la mediana. Si las instituciones europeas siguen midiendo la convergencia con la media aritmética de 37.000 dólares, continuarán legislando de espaldas al 64% de sus naciones miembros.
Segundo: intervenir con urgencia en el corredor oriental para cerrar una brecha inadmisible de 33 a 1.
Y tercero: entender que el bienestar ciudadano no es cuestión de gigantismo industrial, sino de modelo y especialización.

En el plano computacional, todo el análisis es cien por cien reproducible en el script 'lab_sol.py', los gráficos cumplen el estándar de 300 DPI y declaro con total transparencia el uso asistido de Inteligencia Artificial para la verificación y optimización del código.

Europa no necesita más promedios complacientes; necesita mirar de frente a sus medianas.

Muchas gracias por su atención y quedo a su entera disposición para cualquier pregunta." """)

    out_file = "presentacion_demoday1.pptx"
    prs.save(out_file)
    print(f"Clean Presentation (without rubric) saved successfully to {out_file}")

if __name__ == '__main__':
    create_deck()

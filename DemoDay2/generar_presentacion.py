import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def crear_presentacion():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # Diapositiva en blanco

    # Paleta ejecutiva profesional
    C_BG = RGBColor(250, 250, 250)           # #FAFAFA
    C_NAVY = RGBColor(15, 23, 42)            # #0F172A (Texto principal)
    C_BLUE = RGBColor(37, 99, 235)           # #2563EB (Azul corporativo)
    C_BLUE_DARK = RGBColor(30, 58, 138)      # #1E3A8A
    C_GREEN = RGBColor(16, 185, 129)         # #10B981
    C_AMBER = RGBColor(245, 158, 11)         # #F59E0B
    C_MUTED = RGBColor(100, 116, 139)        # #64748B
    C_CARD_BG = RGBColor(255, 255, 255)      # Blanco puro
    C_CARD_BORDER = RGBColor(226, 232, 240)  # #E2E8F0
    C_HIGHLIGHT_BG = RGBColor(241, 245, 249) # #F1F5F9

    BASE_DIR = os.path.dirname(__file__)
    FIG_DIR = os.path.join(BASE_DIR, 'figuras')

    def agregar_fondo(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def agregar_cabecera(slide, categoria, titulo, subtitulo=""):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = categoria.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_BLUE

        tit_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.55))
        tf_tit = tit_box.text_frame
        tf_tit.word_wrap = True
        tf_tit.margin_left = tf_tit.margin_top = tf_tit.margin_right = tf_tit.margin_bottom = 0
        p_tit = tf_tit.paragraphs[0]
        p_tit.text = titulo
        p_tit.font.size = Pt(20)
        p_tit.font.bold = True
        p_tit.font.color.rgb = C_NAVY

        if subtitulo:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.5), Inches(0.4))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitulo
            p_sub.font.size = Pt(11.5)
            p_sub.font.color.rgb = C_MUTED

    # =========================================================================
    # DIAPOSITIVA 1: PORTADA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    agregar_fondo(s1)

    barra = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(1.2), Inches(0.08))
    barra.fill.solid()
    barra.fill.fore_color.rgb = C_BLUE
    barra.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.5), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "Observatorio de Cohesión y Convergencia Europea"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = C_NAVY
    p1.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Diagnóstico Geoespacial y Dinámica Temporal Interactiva (2000–2024)"
    p2.font.size = Pt(18)
    p2.font.color.rgb = C_BLUE_DARK
    p2.space_after = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Cuadro de mando en Plotly Dash nativo con soporte a la decisión estratégica y análisis de convergencia territorial."
    p3.font.size = Pt(13)
    p3.font.color.rgb = C_MUTED

    # Tarjeta de metadatos académicos
    meta_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.5))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = C_CARD_BG
    meta_card.line.color.rgb = C_CARD_BORDER

    tb_meta = s1.shapes.add_textbox(Inches(1.1), Inches(4.4), Inches(11.1), Inches(2.1))
    tf_m = tb_meta.text_frame
    tf_m.word_wrap = True

    pm1 = tf_m.paragraphs[0]
    pm1.text = "DEMO DAY 2: EL SALTO AL DESARROLLO INTERACTIVO (70% EVALUACIÓN)"
    pm1.font.size = Pt(11)
    pm1.font.bold = True
    pm1.font.color.rgb = C_BLUE
    pm1.space_after = Pt(10)

    pm2 = tf_m.add_paragraph()
    pm2.text = "• Asignatura: Visualización y Análisis de Datos (VAD) — Universidad Politécnica de Madrid (UPM)\n" \
               "• Autor: Javier Ferreño\n" \
               "• Tecnologías: Plotly Dash nativo, Folium (Geoespacial HTML), Pandas/GeoPandas y Docker\n" \
               "• Despliegue Público Activo en Render: https://vad-u2iv.onrender.com/"
    pm2.font.size = Pt(12)
    pm2.font.color.rgb = C_NAVY

    # =========================================================================
    # DIAPOSITIVA 2: PREGUNTA ANALÍTICA Y SOPORTE A LA DECISIÓN
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    agregar_fondo(s2)
    agregar_cabecera(s2, "Definición del Problema", "¿Se está cerrando la brecha territorial en Europa?", 
                     "Formulación de la pregunta de negocio y marco de soporte a la decisión de fondos públicos.")

    # 3 Tarjetas de objetivos
    col_w = Inches(3.64)
    gap = Inches(0.38)
    cards_data = [
        ("🎯 PREGUNTA CENTRAL", 
         "¿Se está cerrando la brecha de prosperidad entre el Norte/Occidente y el Sur/Este de Europa?\n\n"
         "El dashboard analiza si la política de cohesión comunitaria ha acelerado la convergencia real o si persisten bolsas estructurales de desigualdad territorial.",
         C_BLUE),
        ("👥 PÚBLICO OBJETIVO", 
         "Planificadores de políticas públicas, comisiones de cohesión de la Unión Europea y directores de fondos estructurales.\n\n"
         "Diseñado para responder a las preguntas prioritarias en menos de 10 segundos con lectura visual inmediata.",
         C_GREEN),
        ("📊 DECISIÓN QUE APOYA", 
         "Asignación focalizada de presupuestos de inversión:\n\n"
         "Priorizar transferencias hacia I+D tecnológica, modernización sanitaria y descarbonización renovable en regiones en transición para evitar trampas de desarrollo.",
         C_AMBER)
    ]

    for i, (tit, desc, col) in enumerate(cards_data):
        x = Inches(0.8) + i * (col_w + gap)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), col_w, Inches(4.8))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = col
        c.line.width = Pt(1.5)

        tb = s2.shapes.add_textbox(x + Inches(0.25), Inches(2.25), col_w - Inches(0.5), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = tit
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = col
        p_t.space_after = Pt(14)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = C_NAVY

    # =========================================================================
    # DIAPOSITIVA 3: ARQUITECTURA TÉCNICA E INTERACTIVIDAD
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    agregar_fondo(s3)
    agregar_cabecera(s3, "Arquitectura y Desarrollo", "Servidor Dash Nativo y Reactividad Dinámica", 
                     "Diseño modular, interactividad continua y controles sincronizados en tiempo real.")

    # Tarjeta izquierda: Características técnicas
    c_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8))
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = C_CARD_BG
    c_left.line.color.rgb = C_CARD_BORDER

    tb_l = s3.shapes.add_textbox(Inches(1.1), Inches(2.25), Inches(5.0), Inches(4.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "PILARES DE LA ARQUITECTURA TÉCNICA"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_BLUE
    p.space_after = Pt(12)

    items_l = [
        ("Plotly Dash Nativo", "Servidor WSGI sobre Flask con componentes Bootswatch Flatly y FontAwesome."),
        ("Análisis Geoespacial en Folium", "Mapas encapsulados en iframe con coropletas continuas y popups HTML interactivos."),
        ("Línea Temporal con Reproductor Automático", "Botones Play/Pausa e intervalo reactivo (1.2s/año) para observar la evolución temporal fluida."),
        ("Callbacks Reactivos en Cascada", "Filtro de macrorregión que actualiza en tiempo real los países disponibles y tarjetas de KPIs."),
        ("Despliegue Multi-Entorno", "Contenerizado en Docker y disponible en vivo en la nube en Render.")
    ]
    for tit, sub in items_l:
        pi = tf_l.add_paragraph()
        pi.text = f"• {tit}: {sub}"
        pi.font.size = Pt(10.5)
        pi.font.color.rgb = C_NAVY
        pi.space_after = Pt(8)

    # Tarjeta derecha: Diagrama de flujo interactivo
    c_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8))
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = C_HIGHLIGHT_BG
    c_right.line.color.rgb = C_BLUE
    c_right.line.width = Pt(1.2)

    tb_r = s3.shapes.add_textbox(Inches(7.1), Inches(2.25), Inches(5.1), Inches(4.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    pr = tf_r.paragraphs[0]
    pr.text = "FLUJO DE REACTIVIDAD (SHNEIDERMAN)"
    pr.font.size = Pt(12)
    pr.font.bold = True
    pr.font.color.rgb = C_BLUE
    pr.space_after = Pt(12)

    flujo_items = [
        "1. Overview First: Panel de filtros globales (Año 2000–2024, Macrorregión, País, Métrica).",
        "2. Zoom & Filter: Al pulsar 'Play', el dashboard avanza año a año actualizando la coropleta y los KPIs superiores.",
        "3. Details-on-Demand: Clic en cualquier país en el mapa para abrir popup HTML con desglose de PIB, longevidad, I+D y desempleo.",
        "4. Enlace Bidireccional: La pestaña temporal (Plotly) vincula el país seleccionado con su serie histórica frente a la media europea."
    ]
    for it in flujo_items:
        pi = tf_r.add_paragraph()
        pi.text = it
        pi.font.size = Pt(11)
        pi.font.color.rgb = C_NAVY
        pi.space_after = Pt(10)

    # =========================================================================
    # DIAPOSITIVA 4: DIAGNÓSTICO ESPACIAL (FOLIUM)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    agregar_fondo(s4)
    agregar_cabecera(s4, "Pestaña 1 · Análisis Espacial", "Distribución Geoespacial y Reducción de Brecha", 
                     "Mapas coropléticos continuos con popups informativos y ranking territorial sincronizado.")

    # Imagen del mapa
    img_map_path = os.path.join(FIG_DIR, 'figura_1_mapa_convergencia.png')
    if os.path.exists(img_map_path):
        s4.shapes.add_picture(img_map_path, Inches(0.8), Inches(2.0), width=Inches(6.8))

    # Tarjeta de hallazgos
    c_hall = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(2.0), Inches(4.6), Inches(4.8))
    c_hall.fill.solid()
    c_hall.fill.fore_color.rgb = C_CARD_BG
    c_hall.line.color.rgb = C_CARD_BORDER

    tb_h = s4.shapes.add_textbox(Inches(8.15), Inches(2.25), Inches(4.1), Inches(4.3))
    tf_h = tb_h.text_frame
    tf_h.word_wrap = True

    ph = tf_h.paragraphs[0]
    ph.text = "EVIDENCIAS GEOESPACIALES CLAVE"
    ph.font.size = Pt(12)
    ph.font.bold = True
    ph.font.color.rgb = C_BLUE
    ph.space_after = Pt(12)

    evidencias = [
        ("Reducción de Brecha Territorial", "La ratio P90/P10 de PIB per cápita se ha reducido de 4.2x (año 2000) a 2.4x (año 2024), demostrando convergencia comunitaria real."),
        ("El Salto de Europa del Este", "Países como Chequia, Polonia y los Bálticos han duplicado su poder adquisitivo medio, acortando distancias con el núcleo occidental."),
        ("Persistencia del Desfase Sur-Norte", "El Sur de Europa experimentó estancamiento post-2008, recuperando ritmo a través de la inversión en transición renovable.")
    ]
    for tit, sub in evidencias:
        pi = tf_h.add_paragraph()
        pi.text = f"• {tit}:\n  {sub}"
        pi.font.size = Pt(11)
        pi.font.color.rgb = C_NAVY
        pi.space_after = Pt(10)

    # =========================================================================
    # DIAPOSITIVA 5: DINÁMICA TEMPORAL & GAPMINDER (PLOTLY)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    agregar_fondo(s5)
    agregar_cabecera(s5, "Pestaña 2 · Dinámica Temporal", "Relación Multivariante Gapminder y Series Históricas", 
                     "Correlación logarítmica entre PIB per cápita, longevidad y tamaño poblacional.")

    # Imagen Gapminder
    img_gap_path = os.path.join(FIG_DIR, 'figura_2_dinamica_gapminder.png')
    if os.path.exists(img_gap_path):
        s5.shapes.add_picture(img_gap_path, Inches(0.8), Inches(2.0), width=Inches(6.8))

    # Tarjeta explicativa
    c_gap = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(2.0), Inches(4.6), Inches(4.8))
    c_gap.fill.solid()
    c_gap.fill.fore_color.rgb = C_CARD_BG
    c_gap.line.color.rgb = C_CARD_BORDER

    tb_g = s5.shapes.add_textbox(Inches(8.15), Inches(2.25), Inches(4.1), Inches(4.3))
    tf_g = tb_g.text_frame
    tf_g.word_wrap = True

    pg = tf_g.paragraphs[0]
    pg.text = "DINÁMICA MULTIVARIANTE"
    pg.font.size = Pt(12)
    pg.font.bold = True
    pg.font.color.rgb = C_GREEN
    pg.space_after = Pt(12)

    g_points = [
        ("Curva de Rendimientos Decrecientes", "Existe una correlación logarítmica fuerte: a partir de $40,000 pc, cada año extra de longevidad exige inversiones exponenciales en I+D sanitaria."),
        ("Convergencia en Esperanza de Vida", "La dispersión en longevidad se ha cerrado mucho más rápido que la de ingresos: todos los países analizados superan ya los 76 años."),
        ("El Dividendo Renovable", "Los países con mayor cuota verde (Dinamarca, Suecia, España, Portugal) muestran mayor estabilidad macroeconómica frente a shocks de precios energéticos.")
    ]
    for tit, sub in g_points:
        pi = tf_g.add_paragraph()
        pi.text = f"• {tit}:\n  {sub}"
        pi.font.size = Pt(11)
        pi.font.color.rgb = C_NAVY
        pi.space_after = Pt(10)

    # =========================================================================
    # DIAPOSITIVA 6: CONCLUSIONES Y SOPORTE A LA DECISIÓN
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    agregar_fondo(s6)
    agregar_cabecera(s6, "Conclusiones y Defensa", "Recomendaciones Estratégicas y Ronda de Preguntas", 
                     "Síntesis ejecutiva de soporte a la decisión para fondos de cohesión comunitarios.")

    col3_w = Inches(3.64)
    conclusiones = [
        ("1. DIAGNÓSTICO", 
         "Convergencia real pero desigual:\n\n"
         "La cohesión europea es una historia de éxito en Europa del Este, pero el Sur aún afronta barreras de productividad y desempleo juvenil.",
         C_BLUE),
        ("2. RECOMENDACIÓN DE FONDOS", 
         "Condicionalidad tecnológica:\n\n"
         "Vincular las partidas comunitarias a alcanzar al menos un 2.5% del PIB en I+D y aceleración renovable para asegurar convergencia antes de 2035.",
         C_GREEN),
        ("3. DEFENSA Y ACCESO", 
         "Aplicación 100% operativa:\n\n"
         "• Despliegue en Render activo\n"
         "• Código modular en lab_sol.py\n"
         "• Gráficos exportados en 300 DPI\n\n"
         "¿Preguntas del tribunal?",
         C_AMBER)
    ]

    for i, (tit, desc, col) in enumerate(conclusiones):
        x = Inches(0.8) + i * (col3_w + gap)
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), col3_w, Inches(4.8))
        c.fill.solid()
        c.fill.fore_color.rgb = C_CARD_BG
        c.line.color.rgb = col
        c.line.width = Pt(1.5)

        tb = s6.shapes.add_textbox(x + Inches(0.25), Inches(2.25), col3_w - Inches(0.5), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = tit
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = col
        p_t.space_after = Pt(14)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = C_NAVY

    out_path = os.path.join(BASE_DIR, 'presentacion_demo_day_2.pptx')
    prs.save(out_path)
    print(f"Presentación guardada con éxito en: {out_path}")

if __name__ == '__main__':
    crear_presentacion()
